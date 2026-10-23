"""Teacher result-entry isolation, validation, review, and publishing tests."""

import unittest
import html
import json
import re
from datetime import datetime, timedelta
from io import BytesIO
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from flask_login import login_required

from app import create_app, db
from app.models import (
    AcademicClass, AcademicLevel, AcademicYear, AcademicYearClass, AcademicYearLevel,
    AcademicSection, AcademicYearSubject, Exam, ExamMarkingConfiguration, GradeScale, Result, Student, StudentEnrollment, Subject, Teacher,
    TeacherActivity, TeacherResultAssignment, TeacherResultEvent, TeacherResultMarkRevision, TeacherResultSubmission, User,
    TeacherResultMark, Setting,
)
from migrations import phase_14_teacher_result_entry
from app.teacher_services import default_teacher_username, validate_teacher_password
from migrations import phase_19_teacher_result_partial
from migrations import phase_20_teacher_result_history
from migrations import phase_21_teacher_result_assignment_archive
from migrations import phase_22_teacher_result_mark_activity
from migrations import phase_24_teacher_result_mark_revisions


class TestTeacherResultEntry(unittest.TestCase):
    class TestConfig:
        TESTING = True
        SECRET_KEY = "teacher-result-entry-test"
        WTF_CSRF_ENABLED = False
        SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
        SQLALCHEMY_TRACK_MODIFICATIONS = False
        AUTO_INIT_DB = False

    def setUp(self):
        self.app = create_app(self.TestConfig)
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.session.remove()
        db.drop_all()
        db.create_all()

        self.year = AcademicYear(name="2026-2027", is_current=True)
        self.level = AcademicLevel(name="Secondary", is_active=True)
        db.session.add_all([self.year, self.level])
        db.session.flush()
        self.year_level = AcademicYearLevel(
            academic_year=self.year, legacy_level=self.level, name="Secondary"
        )
        self.school_class = AcademicClass(academic_level=self.level, name="Form 1", is_active=True)
        db.session.add_all([self.year_level, self.school_class])
        db.session.flush()
        self.year_class = AcademicYearClass(
            academic_year_level=self.year_level, legacy_class=self.school_class, name="Form 1"
        )
        self.subject = Subject(name="Mathematics", academic_level=self.level, max_score=50, is_active=True)
        self.exam = Exam(name="Midterm", academic_year=self.year, is_active=True)
        db.session.add_all([self.year_class, self.subject, self.exam])
        db.session.flush()
        db.session.add(AcademicYearSubject(
            academic_year=self.year, academic_year_level=self.year_level,
            legacy_subject=self.subject, name=self.subject.name, subject_kind="exam", max_score=50,
        ))
        self.students = []
        for index in range(2):
            student = Student(student_code=f"ST-{index + 1}", full_name=f"Student {index + 1}", is_active=True)
            db.session.add(student)
            db.session.flush()
            db.session.add(StudentEnrollment(
                student=student, academic_year=self.year, academic_year_level=self.year_level,
                academic_year_class=self.year_class, status="active", academic_outcome="pending",
                enrollment_source="manual",
            ))
            self.students.append(student)

        self.admin = User(username="result-admin", full_name="Results Admin", role="admin", is_active=True)
        self.admin.set_password("AdminPass123")
        self.admin.set_permissions(["teachers"])
        self.teacher_user = User(username="result-teacher", full_name="Results Teacher", role="staff", is_active=True)
        self.teacher_user.set_password("TeacherPass123")
        db.session.add_all([self.admin, self.teacher_user])
        db.session.flush()
        self.teacher = Teacher(
            teacher_id="RT-1", teacher_code="RT-C1", full_name="Results Teacher",
            employment_status="Active", is_active=True, user_id=self.teacher_user.id,
            classes=[self.school_class], subjects=[self.subject],
        )
        db.session.add(self.teacher)
        db.session.flush()
        self.assignment = TeacherResultAssignment(
            teacher=self.teacher, academic_year=self.year, exam=self.exam,
            school_class=self.school_class, subject=self.subject, max_submissions=1,
        )
        db.session.add(self.assignment)
        db.session.commit()
        self.client = self.app.test_client()

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def _teacher_login(self):
        return self.client.post("/teacher/login", data={
            "username": "result-teacher", "password": "TeacherPass123"
        }, follow_redirects=True)

    def test_teacher_login_prefers_official_school_logo(self):
        db.session.add_all([
            Setting(key="logo_path", value="https://assets.example/school-logo.png"),
            Setting(key="admin_logo_path", value="https://assets.example/admin-logo.png"),
        ])
        db.session.commit()

        response = self.client.get("/teacher/login")

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn('src="https://assets.example/school-logo.png"', body)
        self.assertIn('data-fallback-src="https://assets.example/admin-logo.png"', body)

    def test_teacher_login_username_is_normalized_to_lowercase(self):
        response = self.client.get("/teacher/login")

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn('autocapitalize="none" autocorrect="off" spellcheck="false"', body)
        self.assertIn("username.value.toLowerCase()", body)
        self.assertIn("username.addEventListener('input', normalizeUsername)", body)
        self.assertIn("normalizeUsername();", body)

    def test_teacher_login_includes_social_share_metadata(self):
        self.app.config["PUBLIC_SITE_URL"] = "https://portal.example/"

        response = self.client.get("/teacher/login")

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn('property="og:title" content="Waaxda Macallimiinta |', body)
        self.assertIn('property="og:url" content="https://portal.example/teacher/login"', body)
        self.assertIn('property="og:image" content="https://portal.example/static/images/teacher-portal-share-20261022-v2.jpg"', body)
        self.assertIn('property="og:image:type" content="image/jpeg"', body)
        self.assertIn('property="og:image:width" content="1200"', body)
        self.assertIn('property="og:image:height" content="630"', body)
        self.assertIn('name="twitter:card" content="summary_large_image"', body)
        share_meta = "\n".join(
            line for line in body.splitlines()
            if any(name in line for name in ("name=\"description\"", "og:site_name", "og:description", "og:image:alt", "twitter:description"))
        )
        self.assertNotIn("SULTAAN EMS", share_meta)

        image_response = self.client.get("/static/images/teacher-portal-share-20261022-v2.jpg")
        self.assertEqual(image_response.status_code, 200)
        self.assertEqual(image_response.mimetype, "image/jpeg")
        self.assertGreater(len(image_response.data), 0)

    def test_teacher_back_button_returns_to_login_with_username_only(self):
        portal_response = self._teacher_login()

        self.assertEqual(portal_response.status_code, 200)
        self.assertIn(b'class="icon-btn teacher-back-login"', portal_response.data)
        self.assertIn(b'aria-label="Dib-u-noqo, kuna noqo gelitaanka macallinka"', portal_response.data)

        login_response = self.client.post("/teacher/logout", follow_redirects=True)
        self.assertEqual(login_response.status_code, 200)
        body = login_response.get_data(as_text=True)
        self.assertIn('id="tl-username" name="username" type="text" value="result-teacher"', body)
        password_input = re.search(r'<input id="tl-password"[^>]*>', body)
        self.assertIsNotNone(password_input)
        self.assertNotIn("value=", password_input.group(0))

    def test_result_entry_uses_gendered_student_position_label(self):
        self.students[0].gender = "Male"
        self.students[1].gender = "Female"
        db.session.commit()

        response = self._teacher_login()

        self.assertEqual(response.status_code, 200)
        self.assertIn("Ardeyga 1aad".encode(), response.data)
        self.assertIn("Ardeyadda 2aad".encode(), response.data)

    def test_teacher_password_policy_accepts_letters_or_digits_at_four_characters(self):
        self.assertIsNone(validate_teacher_password("abcd"))
        self.assertIsNone(validate_teacher_password("1234"))
        self.assertIsNone(validate_teacher_password("a1b2"))
        self.assertIsNotNone(validate_teacher_password("abc"))

    def test_default_teacher_username_handles_missing_optional_values(self):
        teacher = SimpleNamespace(teacher_id="TIS-9", teacher_code=None, email=None)
        with patch(
            "app.teacher_services.get_teacher_settings",
            return_value={"teacher_default_username_format": "{teacher_id}.{teacher_code}.{email}"},
        ):
            self.assertEqual(default_teacher_username(teacher), "tis-9..")

    def test_teacher_portal_exposes_only_result_entry_and_limits_roster_to_assignment(self):
        response = self._teacher_login()
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Gelinta natiijooyinka", response.data)
        self.assertIn(b"Student 1", response.data)
        self.assertNotIn(b"class=\"teacher-nav\"", response.data)
        self.assertIn(b"href=\"/teacher/result-entry\"", response.data)
        self.assertIn(b"href=\"/teacher/change-password\"", response.data)
        self.assertIn(b"Beddel password-ka", response.data)
        self.assertNotIn(b"PORTALKA MACALLINKA", response.data)
        self.assertNotIn(b"Exam Analytics", response.data)
        self.assertNotIn(b"Naadigayga", response.data)

    def test_expired_admin_timestamp_does_not_send_teacher_to_admin_login(self):
        with self.client.session_transaction() as session:
            session["admin_last_activity"] = (datetime.utcnow() - timedelta(hours=2)).isoformat()

        response = self._teacher_login()

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Gelinta natiijooyinka", response.data)
        with self.client.session_transaction() as session:
            self.assertNotIn("admin_last_activity", session)

    def test_admin_user_entering_teacher_route_is_sent_to_teacher_login(self):
        self.client.post("/admin/login", data={
            "username": "result-admin", "password": "AdminPass123"
        })

        response = self.client.get("/teacher/result-entry")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/teacher/login", response.headers["Location"])
        self.assertNotIn("/admin", response.headers["Location"])

    def test_anonymous_teacher_page_uses_teacher_login(self):
        response = self.client.get("/teacher/result-entry")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/teacher/login", response.headers["Location"])
        self.assertNotIn("/admin/login", response.headers["Location"])

    def test_flask_login_protected_teacher_route_uses_teacher_login(self):
        def protected_probe():
            return "ok"

        self.app.add_url_rule(
            "/teacher/login-manager-probe",
            endpoint="teacher_login_manager_probe",
            view_func=login_required(protected_probe),
        )

        response = self.client.get("/teacher/login-manager-probe")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/teacher/login", response.headers["Location"])
        self.assertNotIn("/admin/login", response.headers["Location"])

    def test_result_entry_uses_exam_grade_management_and_displays_student_identity(self):
        self.assignment.deadline_at = datetime(2026, 10, 25, 10, 0)
        db.session.commit()
        db.session.add_all([
            ExamMarkingConfiguration(
                academic_year_id=self.year.id,
                academic_year_level_id=self.year_level.id,
                exam_id=self.exam.id,
                default_full_marks=50,
            ),
            GradeScale(
                exam_id=self.exam.id, grade="Exam A", min_score=45, max_score=50,
                comment="Excellent", grade_point=4, is_pass=True,
                badge_color="#ffffff", text_color="#16324f", background_color="#dceeff",
                border_color="#2563eb", sort_order=1, is_active=True,
            ),
            GradeScale(
                exam_id=None, grade="Global fallback", min_score=0, max_score=100,
                comment="Global", grade_point=0, is_pass=False, sort_order=99, is_active=True,
            ),
        ])
        db.session.commit()
        self._teacher_login()

        response = self.client.get("/teacher/result-entry")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"fa-person-dress", response.data)
        self.assertIn(b"Fasalka: Form 1", response.data)
        self.assertIn(b"fa-calendar-day", response.data)
        self.assertIn(b"fa-clock", response.data)
        self.assertIn(b"fa-hourglass-half", response.data)
        self.assertIn(b"fa-stopwatch", response.data)
        self.assertNotIn(b"boqolkiiba ka mid ah 100", response.data)
        markup = response.get_data(as_text=True)
        encoded_scales = re.search(r'data-grade-scales="([^"]*)"', markup)
        self.assertIsNotNone(encoded_scales)
        scales = json.loads(html.unescape(encoded_scales.group(1)))
        self.assertEqual(scales["exam"][0]["grade"], "Exam A")
        self.assertEqual(scales["exam"][0]["min_score"], 90)
        self.assertEqual(scales["exam"][0]["max_score"], 100)
        self.assertEqual(scales["exam"][0]["background_color"], "#dceeff")
        self.assertEqual(scales["exam"][0]["badge_color"], "#ffffff")

    def test_teacher_home_redirects_to_the_real_assigned_result_workspace(self):
        self._teacher_login()
        response = self.client.get("/teacher/?academic_level_id=1")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/teacher/result-entry?academic_level_id=1", response.headers["Location"])

    def test_result_entry_filters_resolve_assignments_across_multiple_levels(self):
        middle_level = AcademicLevel(name="Middle", sort_order=2, is_active=True)
        db.session.add(middle_level)
        db.session.flush()
        middle_year_level = AcademicYearLevel(
            academic_year=self.year, legacy_level=middle_level, name="Middle"
        )
        middle_class = AcademicClass(academic_level=middle_level, name="Grade 7", is_active=True)
        middle_subject = Subject(name="Arabic", academic_level=middle_level, max_score=50, is_active=True)
        db.session.add_all([middle_year_level, middle_class, middle_subject])
        db.session.flush()
        middle_year_class = AcademicYearClass(
            academic_year_level=middle_year_level, legacy_class=middle_class, name="Grade 7"
        )
        db.session.add_all([
            middle_year_class,
            AcademicYearSubject(
                academic_year=self.year, academic_year_level=middle_year_level,
                legacy_subject=middle_subject, name=middle_subject.name,
                subject_kind="exam", max_score=50,
            ),
        ])
        middle_student = Student(student_code="ST-M1", full_name="Middle Student", is_active=True)
        db.session.add(middle_student)
        db.session.flush()
        db.session.add_all([
            StudentEnrollment(
                student=middle_student, academic_year=self.year,
                academic_year_level=middle_year_level, academic_year_class=middle_year_class,
                status="active", academic_outcome="pending", enrollment_source="manual",
            ),
            TeacherResultAssignment(
                teacher=self.teacher, academic_year=self.year, exam=self.exam,
                school_class=middle_class, subject=middle_subject,
            ),
        ])
        db.session.commit()

        self._teacher_login()
        response = self.client.get(
            "/teacher/result-entry",
            query_string={
                "academic_year_id": self.year.id,
                "exam_id": self.exam.id,
                "academic_level_id": middle_level.id,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Middle Student", response.data)
        self.assertIn(b"Grade 7", response.data)
        self.assertIn(b"Arabic", response.data)
        self.assertNotIn(b"Student 1", response.data)
        self.assertNotIn(b"Mathematics", response.data)

    def test_switching_subject_uses_its_own_submission_not_the_previous_sheet(self):
        other_subject = Subject(name="Arabic", academic_level=self.level, max_score=50, is_active=True)
        db.session.add(other_subject)
        db.session.flush()
        db.session.add(AcademicYearSubject(
            academic_year=self.year, academic_year_level=self.year_level,
            legacy_subject=other_subject, name=other_subject.name, subject_kind="exam", max_score=50,
        ))
        other_assignment = TeacherResultAssignment(
            teacher=self.teacher, academic_year=self.year, exam=self.exam,
            school_class=self.school_class, subject=other_subject,
        )
        db.session.add(other_assignment)
        db.session.flush()
        math_submission = TeacherResultSubmission(assignment=self.assignment, status="draft")
        arabic_submission = TeacherResultSubmission(assignment=other_assignment, status="draft")
        db.session.add_all([math_submission, arabic_submission])
        db.session.flush()
        db.session.add_all([
            TeacherResultMark(
                submission=math_submission, student=self.students[0], score=Decimal("42"), status="draft"
            ),
            TeacherResultMark(
                submission=arabic_submission, student=self.students[0], score=Decimal("17"), status="draft"
            ),
        ])
        db.session.commit()
        self._teacher_login()

        response = self.client.get("/teacher/result-entry", query_string={
            "assignment_id": self.assignment.id,
            "academic_year_id": self.year.id,
            "exam_id": self.exam.id,
            "academic_level_id": self.level.id,
            "class_id": self.school_class.id,
            "subject_id": other_subject.id,
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response.headers["Cache-Control"])
        self.assertIn(f'data-assignment-id="{other_assignment.id}"'.encode(), response.data)
        self.assertIn(b"Xaashida hadda furan", response.data)
        self.assertIn(b"<strong>Arabic</strong>", response.data)
        self.assertIn(f'name="score_{self.students[0].id}" value="17.00"'.encode(), response.data)
        self.assertNotIn(f'name="score_{self.students[0].id}" value="42.00"'.encode(), response.data)

    def test_saving_result_keeps_scope_filters_without_sticking_to_stale_status(self):
        self._teacher_login()
        response = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": str(self.assignment.id),
            "action": "draft",
            "academic_year_id": str(self.year.id),
            "exam_id": str(self.exam.id),
            "academic_level_id": str(self.level.id),
            "class_id": str(self.school_class.id),
            "subject_id": str(self.subject.id),
            "status": "not_started",
        })

        self.assertEqual(response.status_code, 302)
        self.assertIn(f"academic_level_id={self.level.id}", response.headers["Location"])
        self.assertIn(f"subject_id={self.subject.id}", response.headers["Location"])
        self.assertNotIn("status=not_started", response.headers["Location"])
        self.assertEqual(TeacherResultSubmission.query.one().status, "draft")

    def test_result_entry_autosave_persists_only_the_selected_assigned_student(self):
        self._teacher_login()
        response = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": self.students[0].id,
            "score": "41,5",
        })

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["success"])
        submission = TeacherResultSubmission.query.one()
        mark = TeacherResultMark.query.one()
        self.assertEqual(submission.status, "draft")
        self.assertEqual(mark.student_id, self.students[0].id)
        self.assertEqual(float(mark.score), 41.5)
        self.assertEqual(response.json["score"], "41.50")

    def test_student_card_and_autosave_track_entry_edit_and_clear_actions(self):
        self._teacher_login()
        payload = {"assignment_id": self.assignment.id, "student_id": self.students[0].id}

        first = self.client.post("/teacher/result-entry/autosave", json={**payload, "score": "31"})
        mark = TeacherResultMark.query.one()
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.json["action_type"], "entered")
        self.assertEqual(mark.last_action, "entered")
        self.assertEqual(mark.last_action_by_id, self.teacher_user.id)

        edited = self.client.post("/teacher/result-entry/autosave", json={**payload, "score": "34"})
        self.assertEqual(edited.json["action_type"], "edited")
        self.assertEqual(mark.last_action, "edited")

        cleared = self.client.post("/teacher/result-entry/autosave", json={**payload, "score": ""})
        self.assertEqual(cleared.json["action_type"], "cleared")
        self.assertEqual(mark.last_action, "cleared")

        page = self.client.get(f"/teacher/result-entry?assignment_id={self.assignment.id}")
        self.assertIn(b"data-action-type=\"cleared\"", page.data)
        self.assertIn("Natiijo waa la tirtiray".encode(), page.data)

    def test_form_save_records_per_student_entry_and_edit(self):
        self._teacher_login()
        first = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "save",
            f"score_{self.students[0].id}": "27",
        }, headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(first.status_code, 200)
        mark = TeacherResultMark.query.one()
        self.assertEqual(mark.last_action, "entered")

        edited = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "save",
            f"score_{self.students[0].id}": "29",
        }, headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(mark.last_action, "edited")
        self.assertEqual(float(mark.score), 29)

    def test_result_entry_autosave_rejects_out_of_range_and_unassigned_students(self):
        self._teacher_login()
        invalid_score = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": self.students[0].id,
            "score": "51",
        })
        too_precise_score = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": self.students[0].id,
            "score": "41.555",
        })
        unassigned_student = Student(student_code="ST-OUT", full_name="Outside Student", is_active=True)
        db.session.add(unassigned_student)
        db.session.commit()
        invalid_student = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": unassigned_student.id,
            "score": "20",
        })

        self.assertEqual(invalid_score.status_code, 422)
        self.assertFalse(invalid_score.json["success"])
        self.assertEqual(too_precise_score.status_code, 422)
        self.assertIn("laba meelood", too_precise_score.json["message"])
        self.assertEqual(invalid_student.status_code, 403)
        self.assertEqual(TeacherResultMark.query.count(), 0)

    def test_result_entry_displays_saved_scores_with_two_decimal_places(self):
        submission = TeacherResultSubmission(assignment=self.assignment)
        db.session.add(submission)
        db.session.flush()
        db.session.add(TeacherResultMark(
            submission=submission, student_id=self.students[0].id,
            score=Decimal("41.236"), status="draft",
        ))
        db.session.commit()
        self._teacher_login()

        response = self.client.get("/teacher/result-entry")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'value="41.24" data-score-input', response.data)
        self.assertNotIn(b'value="41.236" data-score-input', response.data)

    def test_result_entry_json_submit_keeps_partial_submission_behavior(self):
        self._teacher_login()
        response = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": str(self.assignment.id),
            "action": "submit",
            f"score_{self.students[0].id}": "39",
        }, headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"})

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["success"])
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(TeacherResultMark.query.count(), 1)
        self.assertEqual(TeacherResultMark.query.one().student_id, self.students[0].id)

    def test_admin_teacher_form_saves_multi_level_subjects_classes_and_sections(self):
        middle_level = AcademicLevel(name="Middle", sort_order=2, is_active=True)
        legacy_level = AcademicLevel(name="Legacy Primary", sort_order=3, is_active=True)
        db.session.add_all([middle_level, legacy_level])
        db.session.flush()
        middle_year_level = AcademicYearLevel(
            academic_year=self.year, legacy_level=middle_level, name="Dugsi Dhexe", sort_order=2,
        )
        middle_class = AcademicClass(academic_level=middle_level, name="Grade 7", is_active=True)
        middle_subject = Subject(name="Arabic", academic_level=middle_level, is_active=True)
        legacy_class = AcademicClass(academic_level=legacy_level, name="Old Grade", is_active=True)
        legacy_subject = Subject(name="Old Arts", academic_level=legacy_level, is_active=True)
        db.session.add_all([middle_year_level, middle_class, middle_subject, legacy_class, legacy_subject])
        db.session.flush()
        middle_year_class = AcademicYearClass(
            academic_year_level=middle_year_level, legacy_class=middle_class, name="Grade 7",
        )
        middle_year_subject = AcademicYearSubject(
            academic_year=self.year, academic_year_level=middle_year_level,
            legacy_subject=middle_subject, name=middle_subject.name, subject_kind="exam",
        )
        middle_section = AcademicSection(academic_class=middle_class, name="A", is_active=True)
        db.session.add_all([middle_year_class, middle_year_subject, middle_section])
        self.teacher.photo_path = "https://images.example.test/teacher-photo.jpg"
        self.teacher.address = "Teacher address retained"
        self.teacher.emergency_contact = "Emergency contact retained"
        self.teacher.employment_type = "Part Time"
        self.teacher.qualification = "Bachelor degree"
        self.teacher.years_experience = 6
        self.teacher.department = "Science"
        self.teacher.employment_status = "Inactive"
        self.teacher.is_active = False
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        form = self.client.get(f"/admin/teachers/{self.teacher.id}/edit")
        self.assertEqual(form.status_code, 200)
        self.assertIn(b'id="teacher-photo-preview" src="https://images.example.test/teacher-photo.jpg"', form.data)
        self.assertIn(b'data-password-target="teacher-password"', form.data)
        self.assertIn(b'data-password-target="teacher-password-confirm"', form.data)
        self.assertIn(b"data-generate-teacher-password", form.data)
        self.assertIn(b"data-copy-teacher-password", form.data)
        self.assertIn(b"teacher-username-control", form.data)
        self.assertNotIn(b'name="teacher_code"', form.data)
        self.assertNotIn(b'name="address"', form.data)
        self.assertNotIn(b'name="emergency_contact"', form.data)
        self.assertIn(b"Dugsi Dhexe", form.data)
        self.assertIn(b"Arabic", form.data)
        self.assertNotIn(b"Xogta shaqada", form.data)
        self.assertNotIn(b"Taariikhda shaqaalaysiinta", form.data)
        self.assertNotIn(b"Legacy Primary", form.data)
        self.assertNotIn(b"Old Arts", form.data)
        self.assertNotIn(b"Legacy Class Assignments", form.data)

        response = self.client.post(f"/admin/teachers/{self.teacher.id}/edit", data={
            "teacher_id": self.teacher.teacher_id,
            "full_name": self.teacher.full_name,
            "gender": "",
            "phone": "",
            "email": "",
            "academic_level_ids": [str(self.level.id), str(middle_level.id)],
            "academic_class_ids": [str(self.school_class.id), str(middle_class.id)],
            "academic_section_ids": [str(middle_section.id)],
            "subject_ids": [str(self.subject.id), str(middle_subject.id)],
            "username": self.teacher_user.username,
            "account_active": "on",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.teacher)
        self.assertEqual({item.id for item in self.teacher.classes}, {self.school_class.id, middle_class.id})
        self.assertEqual({item.id for item in self.teacher.subjects}, {self.subject.id, middle_subject.id})
        self.assertEqual({item.id for item in self.teacher.sections}, {middle_section.id})
        self.assertEqual({item.id for item in self.teacher.academic_levels}, {self.level.id, middle_level.id})
        self.assertEqual(self.teacher.photo_path, "https://images.example.test/teacher-photo.jpg")
        self.assertEqual(self.teacher.address, "Teacher address retained")
        self.assertEqual(self.teacher.emergency_contact, "Emergency contact retained")
        self.assertEqual(self.teacher.employment_type, "Part Time")
        self.assertEqual(self.teacher.qualification, "Bachelor degree")
        self.assertEqual(self.teacher.years_experience, 6)
        self.assertEqual(self.teacher.department, "Science")
        self.assertEqual(self.teacher.employment_status, "Inactive")
        self.assertFalse(self.teacher.is_active)

    def test_admin_teacher_form_rejects_active_legacy_scope_not_configured_this_year(self):
        legacy_level = AcademicLevel(name="Unconfigured Level", is_active=True)
        db.session.add(legacy_level)
        db.session.flush()
        legacy_class = AcademicClass(academic_level=legacy_level, name="Old Class", is_active=True)
        legacy_subject = Subject(name="Old Subject", academic_level=legacy_level, is_active=True)
        db.session.add_all([legacy_class, legacy_subject])
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post(f"/admin/teachers/{self.teacher.id}/edit", data={
            "teacher_id": self.teacher.teacher_id,
            "teacher_code": self.teacher.teacher_code,
            "full_name": self.teacher.full_name,
            "gender": "",
            "phone": "",
            "email": "",
            "address": "",
            "emergency_contact": "",
            "employment_type": "Full Time",
            "years_experience": "0",
            "department": "",
            "employment_status": "Active",
            "academic_level_ids": [str(self.level.id), str(legacy_level.id)],
            "academic_class_ids": [str(self.school_class.id), str(legacy_class.id)],
            "subject_ids": [str(self.subject.id), str(legacy_subject.id)],
            "username": self.teacher_user.username,
            "account_active": "on",
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn("qaab-dhismeedka sanad-dugsiyeedka hadda".encode(), response.data)
        db.session.refresh(self.teacher)
        self.assertNotIn(legacy_class, self.teacher.classes)
        self.assertNotIn(legacy_subject, self.teacher.subjects)

    def test_admin_password_reset_is_audited_without_exposing_password(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        secret = "NewTeacherPass456"

        response = self.client.post(f"/admin/teachers/{self.teacher.id}/edit", data={
            "teacher_id": self.teacher.teacher_id,
            "teacher_code": self.teacher.teacher_code,
            "full_name": self.teacher.full_name,
            "gender": "",
            "phone": "",
            "email": "",
            "address": "",
            "emergency_contact": "",
            "academic_level_ids": [str(self.level.id)],
            "academic_class_ids": [str(self.school_class.id)],
            "subject_ids": [str(self.subject.id)],
            "username": self.teacher_user.username,
            "password": secret,
            "confirm_password": secret,
            "account_active": "on",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.teacher_user)
        self.assertTrue(self.teacher_user.check_password(secret))
        event = TeacherActivity.query.filter_by(
            teacher_id=self.teacher.id,
            action="Password Reset by Admin",
        ).one()
        self.assertIn("@result-admin", event.details)
        self.assertNotIn(secret, event.details)
        db.session.add(TeacherActivity(
            teacher_id=self.teacher.id,
            action="Password Changed by Teacher",
            details="Macallinku wuxuu beddelay furaha gelitaanka.",
        ))
        db.session.commit()

        profile = self.client.get(f"/admin/teachers/{self.teacher.id}")
        self.assertEqual(profile.status_code, 200)
        self.assertIn(b"Amniga akoonka", profile.data)
        self.assertIn(b"Maamulka ayaa password-ka dib u dejiyey", profile.data)
        self.assertIn(b"Macallinku password-kiisa ayuu beddelay", profile.data)
        self.assertNotIn(secret.encode(), profile.data)

    def test_teacher_changes_password_and_admin_can_review_event(self):
        teacher_client = self.app.test_client()
        response = teacher_client.post("/teacher/login", data={
            "username": self.teacher_user.username,
            "password": "TeacherPass123",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)

        response = teacher_client.post("/teacher/change-password", data={
            "current_password": "TeacherPass123",
            "new_password": "TeacherNew456",
            "confirm_password": "TeacherNew456",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.teacher_user)
        self.assertTrue(self.teacher_user.check_password("TeacherNew456"))
        self.assertEqual(
            TeacherActivity.query.filter_by(
                teacher_id=self.teacher.id,
                action="Password Changed by Teacher",
            ).count(),
            1,
        )


    def test_password_reset_request_requires_matching_teacher_id_and_does_not_change_password(self):
        public_client = self.app.test_client()
        response = public_client.post("/teacher/forgot-password", data={
            "username": self.teacher_user.username,
            "teacher_id": "WRONG-ID",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(TeacherActivity.query.filter_by(
            teacher_id=self.teacher.id,
            action="Password Reset Requested",
        ).count(), 0)

        response = public_client.post("/teacher/forgot-password", data={
            "username": self.teacher_user.username,
            "teacher_id": self.teacher.teacher_id,
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(TeacherActivity.query.filter_by(
            teacher_id=self.teacher.id,
            action="Password Reset Requested",
        ).count(), 1)
        self.assertTrue(self.teacher_user.check_password("TeacherPass123"))

    def test_admin_teacher_directory_filters_academic_class_and_real_portal_access(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.get("/admin/teachers/list", query_string={"class_id": self.school_class.id})
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Results Teacher", response.data)

        self.teacher_user.is_active = False
        db.session.commit()
        disabled = self.client.get("/admin/teachers/list?account_status=disabled")
        enabled = self.client.get("/admin/teachers/list?account_status=enabled")
        self.assertEqual(disabled.status_code, 200)
        self.assertIn(b"Results Teacher", disabled.data)
        self.assertEqual(enabled.status_code, 200)
        self.assertNotIn(b"Results Teacher", enabled.data)

    def test_invalid_score_is_not_saved_and_submitted_marks_publish_only_after_admin_approval(self):
        self._teacher_login()
        invalid = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "draft",
            f"score_{self.students[0].id}": "51",
            f"score_{self.students[1].id}": "40",
        }, follow_redirects=True)
        self.assertEqual(invalid.status_code, 200)
        self.assertIsNone(TeacherResultSubmission.query.first())

        too_precise = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "draft",
            f"score_{self.students[0].id}": "41.555",
        }, headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(too_precise.status_code, 422)
        self.assertIn("laba meelood", too_precise.json["message"])
        self.assertIsNone(TeacherResultSubmission.query.first())

        response = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "49.5",
            f"score_{self.students[1].id}": "40",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(Result.query.count(), 0)

        existing = Result(
            student_id=self.students[0].id, exam_id=self.exam.id,
            subject_id=self.subject.id, score=38, is_published=True,
        )
        db.session.add(existing)
        db.session.commit()

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        review = self.client.get("/admin/teachers/result-entry")
        self.assertEqual(review.status_code, 200)
        approved = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
            "reviewer_note": "Checked against source sheets",
        }, follow_redirects=True)
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 38)
        approved = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
            "reviewer_note": "Checked against source sheets", "confirm_replace": "yes",
        }, follow_redirects=True)
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(submission.status, "approved")
        self.assertEqual(Result.query.count(), 2)
        self.assertTrue(all(result.is_published for result in Result.query.all()))
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 49.5)

    def test_admin_approval_rejects_a_scope_removed_after_teacher_submission(self):
        self._teacher_login()
        submitted = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "42",
        }, follow_redirects=True)
        self.assertEqual(submitted.status_code, 200)
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "submitted")

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        year_subject = AcademicYearSubject.query.filter_by(
            academic_year_id=self.year.id,
            academic_year_level_id=self.year_level.id,
            legacy_subject_id=self.subject.id,
            subject_kind="exam",
        ).one()
        year_subject.is_active = False
        db.session.commit()

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve",
            "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(Result.query.count(), 0)
        self.assertEqual(TeacherResultMark.query.one().status, "submitted")

    def test_teacher_cannot_submit_another_teachers_assignment(self):
        self._teacher_login()
        other_teacher = Teacher(
            teacher_id="RT-2", teacher_code="RT-C2", full_name="Other Teacher",
            employment_status="Active", is_active=True,
        )
        db.session.add(other_teacher)
        db.session.flush()
        other_assignment = TeacherResultAssignment(
            teacher=other_teacher, academic_year=self.year, exam=self.exam,
            school_class=self.school_class, subject=self.subject,
        )
        db.session.add(other_assignment)
        db.session.commit()
        response = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": other_assignment.id, "action": "submit",
        })
        self.assertEqual(response.status_code, 404)
        self.assertEqual(Result.query.count(), 0)

    def test_admin_can_create_assignment_only_for_existing_teacher_scope(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(self.teacher.id),
            "academic_year_id": str(self.year.id), "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(AcademicYearSubject.query.one().id),
            "max_submissions": "2",
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(TeacherResultAssignment.query.count(), 1)
        self.assertEqual(TeacherResultAssignment.query.one().max_submissions, 2)

    def test_admin_cannot_assign_same_class_subject_exam_to_another_teacher(self):
        other_teacher = Teacher(
            teacher_id="RT-2", teacher_code="RT-C2", full_name="Second Teacher",
            employment_status="Active", is_active=True,
        )
        db.session.add(other_teacher)
        db.session.commit()
        year_subject = AcademicYearSubject.query.filter_by(legacy_subject_id=self.subject.id).one()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(other_teacher.id),
            "academic_year_id": str(self.year.id), "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(year_subject.id),
            "max_submissions": "1",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Maaddadan fasalkan iyo imtixaankan macallin kale".encode(), response.data)
        self.assertEqual(TeacherResultAssignment.query.count(), 1)

    def test_admin_scope_picker_receives_existing_class_subject_assignments(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.get("/admin/teachers/result-entry")
        body = response.get_data(as_text=True)
        match = re.search(r'data-reserved-scopes="([^"]*)"', body)

        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(match)
        reservations = json.loads(html.unescape(match.group(1)))
        self.assertEqual(reservations, [{
            "assignment_id": self.assignment.id,
            "academic_year_id": self.year.id,
            "exam_id": self.exam.id,
            "class_id": self.school_class.id,
            "subject_id": self.subject.id,
        }])
        self.assertIn("reservedBySavedAssignment".encode(), response.data)

    def test_admin_new_scope_inherits_last_selected_deadline(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.get("/admin/teachers/result-entry")
        body = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("const inheritedDeadline = rows", body)
        self.assertIn(".find(Boolean) || ''", body)
        self.assertIn(
            'newRow.querySelector(\'[name="deadline_at"]\').value = inheritedDeadline;',
            body,
        )

    def test_admin_scope_picker_orders_exam_before_subject_and_loads_subjects_by_year_level(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.get("/admin/teachers/result-entry")
        body = response.get_data(as_text=True)
        scope_form = body[body.index('data-scope-row'):body.index('<template id="tra-scope-template">')]
        filter_order = [scope_form.index(marker) for marker in (
            "data-year", "data-level", "data-class", "data-exam", "data-subject",
        )]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(filter_order, sorted(filter_order))
        self.assertIn("const subjectsReady = Boolean(yearId && levelId);", body)
        self.assertIn("Choose a class and exam to check which subjects remain available for that scope.", body)

    def test_admin_scope_picker_does_not_guess_level_when_year_has_multiple_levels(self):
        another_level = AcademicLevel(name="Middle", sort_order=2, is_active=True)
        db.session.add(another_level)
        db.session.flush()
        db.session.add(AcademicYearLevel(
            academic_year=self.year, legacy_level=another_level, name="Middle",
        ))
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.get("/admin/teachers/result-entry")
        body = response.get_data(as_text=True)
        level_select = re.search(r'<select name="academic_year_level_id" data-level required>(.*?)</select>', body)

        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(level_select)
        self.assertNotRegex(level_select.group(1), r'<option[^>]*\sselected(?:\s|>)')

    def test_admin_can_save_multiple_teacher_scopes_across_levels_in_one_request(self):
        middle_level = AcademicLevel(name="Middle", sort_order=2, is_active=True)
        db.session.add(middle_level)
        db.session.flush()
        middle_year_level = AcademicYearLevel(
            academic_year=self.year, legacy_level=middle_level, name="Middle"
        )
        middle_class = AcademicClass(academic_level=middle_level, name="Grade 7", is_active=True)
        middle_subject = Subject(name="Arabic", academic_level=middle_level, max_score=50, is_active=True)
        db.session.add_all([middle_year_level, middle_class, middle_subject])
        db.session.flush()
        middle_year_class = AcademicYearClass(
            academic_year_level=middle_year_level, legacy_class=middle_class, name="Grade 7"
        )
        middle_year_subject = AcademicYearSubject(
            academic_year=self.year, academic_year_level=middle_year_level,
            legacy_subject=middle_subject, name=middle_subject.name, subject_kind="exam", max_score=50,
        )
        db.session.add_all([middle_year_class, middle_year_subject])
        db.session.commit()
        math_year_subject = AcademicYearSubject.query.filter_by(legacy_subject_id=self.subject.id).one()

        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(self.teacher.id),
            "academic_year_id": [str(self.year.id), str(self.year.id)],
            "exam_id": [str(self.exam.id), str(self.exam.id)],
            "academic_year_level_id": [str(self.year_level.id), str(middle_year_level.id)],
            "academic_year_class_id": [str(self.year_class.id), str(middle_year_class.id)],
            "academic_year_subject_id": [str(math_year_subject.id), str(middle_year_subject.id)],
            "deadline_at": ["", ""],
            "max_submissions": ["2", "1"],
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        assignments = TeacherResultAssignment.query.filter_by(teacher_id=self.teacher.id).all()
        self.assertEqual(len(assignments), 2)
        by_class = {assignment.class_id: assignment for assignment in assignments}
        self.assertEqual(by_class[self.school_class.id].max_submissions, 2)
        self.assertEqual(by_class[middle_class.id].subject_id, middle_subject.id)
        self.assertIn(b"1 xilsaaris cusub iyo 1 xilsaaris oo la cusboonaysiiyey", response.data)

    def test_admin_multi_scope_creation_is_atomic_when_any_scope_is_invalid(self):
        middle_level = AcademicLevel(name="Middle", is_active=True)
        db.session.add(middle_level)
        db.session.flush()
        middle_year_level = AcademicYearLevel(academic_year=self.year, legacy_level=middle_level, name="Middle")
        middle_class = AcademicClass(academic_level=middle_level, name="Grade 7", is_active=True)
        middle_subject = Subject(name="Arabic", academic_level=middle_level, is_active=True)
        db.session.add_all([middle_year_level, middle_class, middle_subject])
        db.session.flush()
        middle_year_class = AcademicYearClass(academic_year_level=middle_year_level, legacy_class=middle_class, name="Grade 7")
        middle_year_subject = AcademicYearSubject(
            academic_year=self.year, academic_year_level=middle_year_level,
            legacy_subject=middle_subject, name=middle_subject.name, subject_kind="exam",
        )
        db.session.add_all([middle_year_class, middle_year_subject])
        db.session.commit()
        math_year_subject = AcademicYearSubject.query.filter_by(legacy_subject_id=self.subject.id).one()

        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(self.teacher.id),
            "academic_year_id": [str(self.year.id), str(self.year.id)],
            "exam_id": [str(self.exam.id), str(self.exam.id)],
            "academic_year_level_id": [str(self.year_level.id), str(middle_year_level.id)],
            "academic_year_class_id": [str(self.year_class.id), str(self.year_class.id)],
            "academic_year_subject_id": [str(math_year_subject.id), str(middle_year_subject.id)],
            "max_submissions": ["3", "1"],
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Scope kasta waa inuu si sax ah".encode(), response.data)
        self.assertEqual(TeacherResultAssignment.query.count(), 1)
        self.assertEqual(self.assignment.max_submissions, 1)

    def test_admin_can_assign_result_scope_without_legacy_teacher_class_subject_links(self):
        unscoped_subject = Subject(
            name="Science", academic_level=self.level, max_score=50, is_active=True,
        )
        unscoped_teacher = Teacher(
            teacher_id="RT-3", teacher_code="RT-C3", full_name="Unscoped Teacher",
            employment_status="Active", is_active=True,
        )
        db.session.add_all([unscoped_teacher, unscoped_subject])
        db.session.flush()
        unscoped_year_subject = AcademicYearSubject(
            academic_year=self.year, academic_year_level=self.year_level,
            legacy_subject=unscoped_subject, name=unscoped_subject.name,
            subject_kind="exam", max_score=50,
        )
        db.session.add(unscoped_year_subject)
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(unscoped_teacher.id),
            "academic_year_id": str(self.year.id), "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(unscoped_year_subject.id),
            "max_submissions": "1",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        assignment = TeacherResultAssignment.query.filter_by(teacher_id=unscoped_teacher.id).one()
        self.assertEqual(assignment.class_id, self.school_class.id)
        self.assertEqual(assignment.subject_id, unscoped_subject.id)

    def test_admin_assignment_form_uses_year_scoped_records(self):
        orphan_class = AcademicClass(academic_level=self.level, name="Unmapped legacy class", is_active=True)
        orphan_subject = Subject(name="Unoffered legacy subject", academic_level=self.level, is_active=True)
        db.session.add_all([orphan_class, orphan_subject])
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.get("/admin/teachers/result-entry")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"name=\"academic_year_level_id\"", response.data)
        self.assertIn(b"name=\"academic_year_class_id\"", response.data)
        self.assertIn(b"name=\"academic_year_subject_id\"", response.data)
        self.assertIn(f'data-level-id="{self.year_level.id}"'.encode(), response.data)
        self.assertIn(b"Mathematics", response.data)
        self.assertIn(b"tra-add-scope", response.data)
        self.assertIn(b"data-scope-row", response.data)
        self.assertNotIn(b"Unmapped legacy class", response.data)
        self.assertNotIn(b"Unoffered legacy subject", response.data)

    def test_admin_can_edit_assignment_scope_and_policy_before_submission(self):
        other_subject = Subject(
            name="Arabic", academic_level=self.level, max_score=50, is_active=True
        )
        db.session.add(other_subject)
        db.session.flush()
        other_year_subject = AcademicYearSubject(
            academic_year=self.year,
            academic_year_level=self.year_level,
            legacy_subject=other_subject,
            name=other_subject.name,
            subject_kind="exam",
            max_score=50,
        )
        db.session.add(other_year_subject)
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        page = self.client.get("/admin/teachers/result-entry")
        self.assertIn(b'class="tra-edit-trigger"', page.data)
        self.assertIn(f'data-assignment-id="{self.assignment.id}"'.encode(), page.data)

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "update",
            "assignment_id": str(self.assignment.id),
            "teacher_id": str(self.teacher.id),
            "academic_year_id": str(self.year.id),
            "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(other_year_subject.id),
            "deadline_at": "2026-10-22T10:00",
            "max_submissions": "3",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.assignment)
        self.assertEqual(self.assignment.subject_id, other_subject.id)
        self.assertEqual(self.assignment.max_submissions, 3)
        self.assertEqual(self.assignment.deadline_at.strftime("%Y-%m-%d %H:%M"), "2026-10-22 07:00")
        self.assertEqual(TeacherResultAssignment.query.count(), 1)
        self.assertEqual(TeacherResultEvent.query.filter_by(
            assignment_id=self.assignment.id, event_type="assignment_updated"
        ).count(), 1)
        self.assertIn("Xilsaarista waa la tafatiray".encode(), response.data)

    def test_admin_can_update_policy_on_legacy_duplicate_scope(self):
        other_teacher = Teacher(
            teacher_id="RT-4", teacher_code="RT-C4", full_name="Legacy Duplicate Teacher",
            employment_status="Active", is_active=True,
        )
        db.session.add(other_teacher)
        db.session.flush()
        db.session.add(TeacherResultAssignment(
            teacher=other_teacher, academic_year=self.year, exam=self.exam,
            school_class=self.school_class, subject=self.subject,
        ))
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "update", "assignment_id": str(self.assignment.id),
            "teacher_id": str(self.teacher.id),
            "academic_year_id": str(self.year.id), "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(AcademicYearSubject.query.one().id),
            "deadline_at": "2026-10-23T10:00", "max_submissions": "2",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.assignment)
        self.assertEqual(self.assignment.max_submissions, 2)
        self.assertEqual(self.assignment.deadline_at.strftime("%Y-%m-%d %H:%M"), "2026-10-23 07:00")

    def test_admin_cannot_change_assignment_scope_after_submission_exists(self):
        other_subject = Subject(
            name="Arabic", academic_level=self.level, max_score=50, is_active=True
        )
        db.session.add(other_subject)
        db.session.flush()
        other_year_subject = AcademicYearSubject(
            academic_year=self.year,
            academic_year_level=self.year_level,
            legacy_subject=other_subject,
            name=other_subject.name,
            subject_kind="exam",
            max_score=50,
        )
        submission = TeacherResultSubmission(assignment=self.assignment, status="draft")
        db.session.add_all([other_year_subject, submission])
        db.session.flush()
        db.session.add(TeacherResultMark(
            submission=submission, student=self.students[0], score=Decimal("42"), status="draft"
        ))
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "update",
            "assignment_id": str(self.assignment.id),
            "teacher_id": str(self.teacher.id),
            "academic_year_id": str(self.year.id),
            "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(self.year_class.id),
            "academic_year_subject_id": str(other_year_subject.id),
            "deadline_at": "",
            "max_submissions": "2",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        db.session.refresh(self.assignment)
        self.assertEqual(self.assignment.subject_id, self.subject.id)
        self.assertIn("xaashida natiijadu way jirtaa".encode(), response.data.lower())
        db.session.refresh(submission)
        self.assertEqual(submission.marks[0].score, Decimal("42.000"))

    def test_unlinked_staff_account_can_be_linked_and_then_used_for_result_assignment(self):
        portal_user = User(
            username="new-teacher-login", full_name="New Teacher", role="staff", is_active=True,
        )
        portal_user.set_password("TeacherPass123")
        db.session.add(portal_user)
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        setup = self.client.get("/admin/teachers/result-entry")
        self.assertEqual(setup.status_code, 200)
        self.assertIn(b"new-teacher-login", setup.data)
        self.assertIn(f"user_id={portal_user.id}".encode(), setup.data)

        form = self.client.get(f"/admin/teachers/new?user_id={portal_user.id}")
        self.assertEqual(form.status_code, 200)
        self.assertIn(b"new-teacher-login", form.data)
        response = self.client.post("/admin/teachers/new", data={
            "existing_user_id": str(portal_user.id),
            "teacher_id": "NEW-TEACHER",
            "teacher_code": "",
            "full_name": "New Teacher",
            "employment_type": "Full Time",
            "years_experience": "0",
            "employment_status": "Active",
            "academic_level_ids": [str(self.level.id)],
            "academic_class_ids": [str(self.school_class.id)],
            "subject_ids": [str(self.subject.id)],
            "username": portal_user.username,
            "account_active": "on",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        linked_teacher = Teacher.query.filter_by(teacher_id="NEW-TEACHER").one()
        db.session.refresh(portal_user)
        self.assertEqual(linked_teacher.user_id, portal_user.id)
        self.assertEqual(portal_user.role, "staff")
        self.assertEqual(Teacher.query.filter_by(user_id=portal_user.id).count(), 1)

        setup = self.client.get("/admin/teachers/result-entry")
        self.assertIn(f'value="{linked_teacher.id}">{linked_teacher.full_name}'.encode(), setup.data)
        self.assertNotIn(b"new-teacher-login", setup.data)

    def test_teacher_form_refuses_to_link_an_admin_login_as_a_teacher(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.post("/admin/teachers/new", data={
            "existing_user_id": str(self.admin.id),
            "teacher_id": "INVALID-TEACHER",
            "full_name": "Should Not Link",
            "employment_type": "Full Time",
            "years_experience": "0",
            "employment_status": "Active",
            "academic_level_ids": [str(self.level.id)],
            "academic_class_ids": [str(self.school_class.id)],
            "subject_ids": [str(self.subject.id)],
            "username": self.admin.username,
            "account_active": "on",
        })

        self.assertEqual(response.status_code, 200)
        self.assertIn("Akoonka la doortay ma aha staff".encode(), response.data)
        self.assertIsNone(Teacher.query.filter_by(teacher_id="INVALID-TEACHER").first())
        db.session.refresh(self.admin)
        self.assertEqual(self.admin.role, "admin")

    def test_admin_rejects_class_and_subject_scopes_from_another_level(self):
        other_level = AcademicLevel(name="Middle", is_active=True)
        db.session.add(other_level)
        db.session.flush()
        other_year_level = AcademicYearLevel(
            academic_year=self.year, legacy_level=other_level, name="Middle"
        )
        other_class = AcademicClass(academic_level=other_level, name="Grade 7", is_active=True)
        other_subject = Subject(name="Science", academic_level=other_level, is_active=True)
        db.session.add_all([other_year_level, other_class, other_subject])
        db.session.flush()
        other_year_class = AcademicYearClass(
            academic_year_level=other_year_level, legacy_class=other_class, name="Grade 7"
        )
        other_year_subject = AcademicYearSubject(
            academic_year=self.year, academic_year_level=other_year_level,
            legacy_subject=other_subject, name=other_subject.name, subject_kind="exam",
        )
        db.session.add_all([other_year_class, other_year_subject])
        db.session.flush()
        self.teacher.classes.append(other_class)
        self.teacher.subjects.append(other_subject)
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        response = self.client.post("/admin/teachers/result-entry", data={
            "action": "create", "teacher_id": str(self.teacher.id),
            "academic_year_id": str(self.year.id), "exam_id": str(self.exam.id),
            "academic_year_level_id": str(self.year_level.id),
            "academic_year_class_id": str(other_year_class.id),
            "academic_year_subject_id": str(other_year_subject.id),
            "max_submissions": "1",
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Scope kasta waa inuu si sax ah".encode(), response.data)
        self.assertEqual(TeacherResultAssignment.query.count(), 1)

    def test_admin_teacher_home_opens_result_entry_control_without_dashboard_mockups(self):
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        response = self.client.get("/admin/teachers/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/teachers/result-entry", response.headers["Location"])

        response = self.client.get("/admin/teachers/result-entry")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Result Entry Control", response.data)
        self.assertIn(b"All assignments", response.data)
        self.assertIn(b"Form 1", response.data)
        self.assertNotIn(b"Top Performance Snapshot", response.data)

    def test_reopened_approved_sheet_requires_reason_and_explicit_result_replacement(self):
        self.client.post("/teacher/login", data={
            "username": "result-teacher", "password": "TeacherPass123"
        })
        first_submit = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "44",
            f"score_{self.students[1].id}": "39",
        }, follow_redirects=True)
        self.assertEqual(first_submit.status_code, 200)

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "approved")

        missing_reason = self.client.post("/admin/teachers/result-entry", data={
            "action": "reopen", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(missing_reason.status_code, 200)
        self.assertIn(b'tra-reason-error', missing_reason.data)
        self.assertIn("Fadlan qor sababta dib loogu furayo ka hor intaadan kaydin.".encode(), missing_reason.data)
        self.assertEqual(submission.status, "approved")
        reopened = self.client.post("/admin/teachers/result-entry", data={
            "action": "reopen", "assignment_id": str(self.assignment.id),
            "reviewer_note": "Correct one entry after source-sheet review",
        }, follow_redirects=True)
        self.assertEqual(reopened.status_code, 200)
        self.assertEqual(submission.status, "reopened")
        self.assertEqual(self.assignment.max_submissions, 2)

        self.client.post("/admin/logout")
        self.client.post("/teacher/login", data={
            "username": "result-teacher", "password": "TeacherPass123"
        })
        cleared = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": self.students[0].id,
            "score": "",
        })
        self.assertEqual(cleared.status_code, 200)
        self.assertIsNone(cleared.get_json()["score"])
        reopened_page = self.client.get(f"/teacher/result-entry?assignment_id={self.assignment.id}")
        self.assertIn(f'name="score_{self.students[0].id}" value=""'.encode(), reopened_page.data)
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 44)
        second_submit = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "46",
            f"score_{self.students[1].id}": "39",
        }, follow_redirects=True)
        self.assertEqual(second_submit.status_code, 200)
        self.assertEqual(submission.submission_count, 2)

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 44)
        self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id), "confirm_replace": "yes",
        }, follow_redirects=True)
        self.assertEqual(submission.status, "approved")
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 46)

    def test_admin_can_return_submitted_marks_for_correction_then_approve_resubmission(self):
        self._teacher_login()
        self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "44",
            f"score_{self.students[1].id}": "39",
        }, follow_redirects=True)
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "submitted")
        first_revisions = TeacherResultMarkRevision.query.order_by(TeacherResultMarkRevision.id).all()
        self.assertEqual(len(first_revisions), 2)
        self.assertTrue(all(not revision.is_change for revision in first_revisions))
        self.assertTrue(all(revision.entry_status == "submitted" for revision in first_revisions))
        self.assertTrue(all(revision.submission_attempt == 1 for revision in first_revisions))

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        first_review = self.client.get("/admin/teachers/result-entry?status=review")
        self.assertEqual(first_review.status_code, 200)
        self.assertIn(b"Gelintii bilowga", first_review.data)
        self.assertIn(b"Gudbin 1", first_review.data)
        missing_reason = self.client.post("/admin/teachers/result-entry", data={
            "action": "reopen", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertIn("Qor sababta dib loogu furayo".encode(), missing_reason.data)
        self.assertIn(b'tra-reason-error', missing_reason.data)
        self.assertIn("Fadlan qor sababta dib loogu furayo ka hor intaadan kaydin.".encode(), missing_reason.data)
        self.assertEqual(submission.status, "submitted")

        reopened = self.client.post("/admin/teachers/result-entry", data={
            "action": "reopen", "assignment_id": str(self.assignment.id),
            "reviewer_note": "Corrected against the signed source sheet",
        }, follow_redirects=True)
        self.assertEqual(reopened.status_code, 200)
        self.assertEqual(submission.status, "reopened")
        self.assertTrue(all(mark.status == "draft" for mark in submission.marks))

        self.client.post("/admin/logout")
        self._teacher_login()
        correction = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "46.25",
            f"score_{self.students[1].id}": "39",
        }, follow_redirects=True)
        self.assertEqual(correction.status_code, 200)
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(submission.submission_count, 2)
        first_mark = TeacherResultMark.query.filter_by(student_id=self.students[0].id).one()
        first_student_revisions = TeacherResultMarkRevision.query.filter_by(mark_id=first_mark.id).order_by(
            TeacherResultMarkRevision.id
        ).all()
        self.assertEqual(len(first_student_revisions), 2)
        self.assertEqual(first_student_revisions[1].previous_score, Decimal("44.00"))
        self.assertEqual(first_student_revisions[1].next_score, Decimal("46.25"))
        self.assertEqual(first_student_revisions[1].action, "edited")
        self.assertTrue(first_student_revisions[1].is_change)
        self.assertEqual(first_student_revisions[1].submission_attempt, 2)
        second_mark = TeacherResultMark.query.filter_by(student_id=self.students[1].id).one()
        self.assertEqual(TeacherResultMarkRevision.query.filter_by(mark_id=second_mark.id).count(), 1)

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        second_review = self.client.get("/admin/teachers/result-entry?status=review")
        self.assertEqual(second_review.status_code, 200)
        self.assertIn(b"1 isbeddel", second_review.data)
        self.assertIn(b"Isbeddelka 1aad", second_review.data)
        self.assertIn(b"Gudbin 2", second_review.data)
        approved = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(submission.status, "approved")
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 46.25)
        approved_history = self.client.get("/admin/teachers/result-entry?status=approved")
        self.assertEqual(approved_history.status_code, 200)
        self.assertIn(b"Taariikhda dhibcaha", approved_history.data)
        self.assertIn(b"46.25", approved_history.data)
        self.assertIn(b"44.00", approved_history.data)

    def test_partial_submission_approval_publishes_only_entered_students_and_allows_continuation(self):
        self.students[0].photo_path = "uploads/students/student-1.jpg"
        self.students[0].mother_name = "Sahra Said"
        db.session.commit()
        self._teacher_login()
        first = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "44",
            f"score_{self.students[1].id}": "",
        }, follow_redirects=True)

        self.assertEqual(first.status_code, 200)
        submission = TeacherResultSubmission.query.one()
        self.assertEqual(submission.status, "submitted")
        self.assertEqual(submission.submission_count, 1)
        first_mark = TeacherResultMark.query.filter_by(student_id=self.students[0].id).one()
        self.assertEqual(first_mark.status, "submitted")
        self.assertEqual(Result.query.count(), 0)

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        review = self.client.get("/admin/teachers/result-entry?status=review")
        self.assertEqual(review.status_code, 200)
        self.assertIn(b"Review 1 submitted student marks", review.data)
        self.assertIn(b"/static/uploads/students/student-1.jpg", review.data)
        self.assertIn(b"Student 1", review.data)
        self.assertIn(b"ST-1", review.data)
        self.assertIn(b"Sahra Said", review.data)
        self.assertIn(b"fa-person-dress", review.data)
        self.assertIn(b"Natiijada rasmiga ah", review.data)
        self.assertNotIn(b"Student 2</td>", review.data)
        approved_first = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(approved_first.status_code, 200)
        self.assertEqual(submission.status, "draft")
        self.assertEqual(first_mark.status, "approved")
        self.assertEqual(float(Result.query.filter_by(student_id=self.students[0].id).one().score), 44)
        self.assertIsNone(Result.query.filter_by(student_id=self.students[1].id).first())

        self.client.post("/admin/logout")
        self._teacher_login()
        workspace = self.client.get(f"/teacher/result-entry?assignment_id={self.assignment.id}")
        self.assertEqual(workspace.status_code, 200)
        self.assertIn(b"Qaybo la ansixiyey", workspace.data)
        self.assertIn(f'name="score_{self.students[0].id}"'.encode(), workspace.data)
        self.assertIn(f'name="score_{self.students[1].id}"'.encode(), workspace.data)
        self.assertIn(f'name="score_{self.students[0].id}" value="44.00"'.encode(), workspace.data)
        self.assertIn(f'name="score_{self.students[1].id}" value=""'.encode(), workspace.data)
        self.assertRegex(workspace.data, f'name="score_{self.students[0].id}"[^>]*disabled'.encode())
        self.assertNotRegex(workspace.data, f'name="score_{self.students[1].id}"[^>]*disabled'.encode())

        second = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "44",
            f"score_{self.students[1].id}": "39",
        }, follow_redirects=True)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(submission.submission_count, 1)
        second_mark = TeacherResultMark.query.filter_by(student_id=self.students[1].id).one()
        self.assertEqual(second_mark.status, "submitted")

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        approved_second = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(approved_second.status_code, 200)
        self.assertEqual(submission.status, "approved")
        self.assertEqual(second_mark.status, "approved")
        self.assertEqual(Result.query.count(), 2)
        self.assertEqual(submission.submission_count, 1)

    def test_result_entry_history_tracks_review_events_filters_and_exports(self):
        self._teacher_login()
        submitted = self.client.post("/teacher/result-entry/save", data={
            "assignment_id": self.assignment.id,
            "action": "submit",
            f"score_{self.students[0].id}": "42",
        }, follow_redirects=True)
        self.assertEqual(submitted.status_code, 200)
        first_event = TeacherResultEvent.query.filter_by(event_type="submitted").one()
        self.assertEqual(first_event.marks_count, 1)
        self.assertEqual(first_event.teacher_name, "Results Teacher")

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        approved = self.client.post("/admin/teachers/result-entry", data={
            "action": "approve", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(approved.status_code, 200)
        approval_event = TeacherResultEvent.query.filter_by(event_type="partially_approved").one()
        self.assertEqual(approval_event.marks_count, 1)

        history = self.client.get("/admin/teachers/result-entry/history")
        self.assertEqual(history.status_code, 200)
        self.assertIn("Taariikhda Gelinta Natiijooyinka".encode(), history.data)
        self.assertIn("Natiijo loo gudbiyey maamulka".encode(), history.data)
        self.assertIn("Qayb ahaan la ansixiyey".encode(), history.data)

        filtered = self.client.get("/admin/teachers/result-entry/history?event_type=partially_approved&teacher=Results%20Teacher")
        self.assertEqual(filtered.status_code, 200)
        self.assertIn("Qayb ahaan la ansixiyey".encode(), filtered.data)
        self.assertNotIn(b'class="trh-event submitted"', filtered.data)

        report = self.client.get("/admin/teachers/result-entry/history/export.xlsx?event_type=partially_approved")
        self.assertEqual(report.status_code, 200)
        sheet = load_workbook(BytesIO(report.data), read_only=True).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 2).value, "Qayb ahaan la ansixiyey")

    def test_admin_can_archive_restore_assignment_and_history_keeps_scope_and_marks(self):
        submission = TeacherResultSubmission(assignment=self.assignment, status="draft")
        db.session.add(submission)
        db.session.flush()
        mark = TeacherResultMark(
            submission=submission, student=self.students[0], score=Decimal("42"), status="draft"
        )
        db.session.add(mark)
        db.session.commit()

        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        archived = self.client.post("/admin/teachers/result-entry", data={
            "action": "delete", "assignment_id": str(self.assignment.id),
            "deletion_note": "Xilsaaris khaldan ayaa la saxayaa.",
        }, follow_redirects=True)
        self.assertEqual(archived.status_code, 200)
        self.assertIn("La tirtiray".encode(), archived.data)
        self.assertIsNotNone(self.assignment.deleted_at)
        self.assertFalse(self.assignment.is_active)
        self.assertEqual(TeacherResultSubmission.query.count(), 1)
        self.assertEqual(TeacherResultMark.query.one().score, Decimal("42"))

        deleted_event = TeacherResultEvent.query.filter_by(event_type="assignment_deleted").one()
        self.assertEqual(deleted_event.teacher_name, "Results Teacher")
        self.assertEqual(deleted_event.class_name, "Form 1")
        self.assertEqual(deleted_event.subject_name, "Mathematics")
        self.assertEqual(deleted_event.status_from, "active")
        self.assertEqual(deleted_event.status_to, "deleted")
        self.assertEqual(deleted_event.note, "Xilsaaris khaldan ayaa la saxayaa.")

        deleted_list = self.client.get("/admin/teachers/result-entry?status=deleted")
        self.assertEqual(deleted_list.status_code, 200)
        self.assertIn("Dib u soo celi".encode(), deleted_list.data)
        self.assertIn("xogta iyo dhibcihii hore waa la hayaa".encode(), deleted_list.data)

        self.client.post("/admin/logout")
        self._teacher_login()
        blocked = self.client.post("/teacher/result-entry/autosave", json={
            "assignment_id": self.assignment.id,
            "student_id": self.students[0].id,
            "score": "49",
        })
        self.assertEqual(blocked.status_code, 404)

        self.client.post("/teacher/logout")
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})
        history = self.client.get("/admin/teachers/result-entry/history?event_type=assignment_deleted")
        self.assertEqual(history.status_code, 200)
        self.assertIn("Xilsaaris kayd ahaan loo tirtiray".encode(), history.data)
        self.assertIn("Xilsaaris khaldan ayaa la saxayaa.".encode(), history.data)

        restored = self.client.post("/admin/teachers/result-entry", data={
            "action": "restore", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(restored.status_code, 200)
        self.assertIsNone(self.assignment.deleted_at)
        self.assertTrue(self.assignment.is_active)
        self.assertEqual(TeacherResultSubmission.query.count(), 1)
        self.assertEqual(TeacherResultMark.query.one().score, Decimal("42"))
        self.assertEqual(TeacherResultEvent.query.filter_by(event_type="assignment_restored").count(), 1)

    def test_admin_can_archive_unstarted_assignment_and_restore_its_paused_state(self):
        self.assignment.is_active = False
        db.session.commit()
        self.client.post("/admin/login", data={"username": "result-admin", "password": "AdminPass123"})

        archived = self.client.post("/admin/teachers/result-entry", data={
            "action": "delete", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(archived.status_code, 200)
        self.assertEqual(TeacherResultSubmission.query.count(), 0)
        self.assertIsNone(TeacherResultEvent.query.filter_by(event_type="assignment_deleted").one().submission_id)

        restored = self.client.post("/admin/teachers/result-entry", data={
            "action": "restore", "assignment_id": str(self.assignment.id),
        }, follow_redirects=True)
        self.assertEqual(restored.status_code, 200)
        self.assertIsNone(self.assignment.deleted_at)
        self.assertFalse(self.assignment.is_active)
        event = TeacherResultEvent.query.filter_by(event_type="assignment_restored").one()
        self.assertEqual(event.status_from, "deleted")
        self.assertEqual(event.status_to, "paused")

    def test_migration_is_repeatable(self):
        with patch.object(phase_14_teacher_result_entry, "create_app", return_value=self.app):
            self.assertEqual(phase_14_teacher_result_entry.upgrade(), {"status": "applied", "tables": 3})
            self.assertEqual(phase_14_teacher_result_entry.upgrade(), {"status": "applied", "tables": 3})
        with patch.object(phase_19_teacher_result_partial, "create_app", return_value=self.app):
            self.assertEqual(phase_19_teacher_result_partial.upgrade(), {"status": "applied", "columns": 1})
            self.assertEqual(phase_19_teacher_result_partial.upgrade(), {"status": "applied", "columns": 1})
        db.session.add(TeacherResultSubmission(assignment=self.assignment, status="draft"))
        db.session.commit()
        with patch.object(phase_20_teacher_result_history, "create_app", return_value=self.app):
            self.assertEqual(phase_20_teacher_result_history.upgrade(), {"status": "applied", "table": "teacher_result_events"})
            baseline_count = TeacherResultEvent.query.filter_by(event_type="baseline").count()
            self.assertEqual(baseline_count, 1)
            self.assertEqual(phase_20_teacher_result_history.upgrade(), {"status": "applied", "table": "teacher_result_events"})
            self.assertEqual(TeacherResultEvent.query.filter_by(event_type="baseline").count(), baseline_count)
        with patch.object(phase_21_teacher_result_assignment_archive, "create_app", return_value=self.app):
            self.assertEqual(phase_21_teacher_result_assignment_archive.upgrade(), {"status": "applied", "columns": 2})
            self.assertEqual(phase_21_teacher_result_assignment_archive.upgrade(), {"status": "applied", "columns": 2})
        with patch.object(phase_22_teacher_result_mark_activity, "create_app", return_value=self.app):
            self.assertEqual(phase_22_teacher_result_mark_activity.upgrade(), {"status": "applied", "columns": 3})
            self.assertEqual(phase_22_teacher_result_mark_activity.upgrade(), {"status": "applied", "columns": 3})
        with patch.object(phase_24_teacher_result_mark_revisions, "create_app", return_value=self.app):
            expected = {"status": "applied", "table": "teacher_result_mark_revisions"}
            self.assertEqual(phase_24_teacher_result_mark_revisions.upgrade(), expected)
            self.assertEqual(phase_24_teacher_result_mark_revisions.upgrade(), expected)


if __name__ == "__main__":
    unittest.main()
